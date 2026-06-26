"""Time Clock Bulk Upload — user entry tool.

A small standalone desktop app for building a bulk-upload file for a time
clock system.

What it does
------------
* Lets you enter as many users as you need (ID, First Name, Last Name,
  Pin, Passcode).
* Only writes the upload file once you click *Generate File*.
* Writes ``Ecuser.txt`` where every user is one line and every field is
  written as ``fieldID=fieldvalue`` with fields separated by commas, e.g.::

      id=1,First=John,Last=Doe,Pin=1234,Passcode=5678

* Can be linked to an Excel workbook that keeps a permanent record of every
  user ever added. The workbook is read on link so IDs that already exist
  are rejected as duplicates, and it is appended to every time you generate
  a new batch.

Field IDs written to the file:
    ID        -> "id"
    First Name-> "First"
    Last Name -> "Last"
    Pin       -> "Pin"
    Passcode  -> "Passcode"

Build a Windows .exe with PyInstaller (see README.md):
    pyinstaller --onefile --windowed --name TimeClockBulkUpload time_clock_bulk_upload.py
"""

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# openpyxl is optional. The app still works without it (the Excel record
# feature is simply disabled and we tell the user).
try:
    from openpyxl import Workbook, load_workbook

    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover - depends on environment
    HAVE_OPENPYXL = False


# Field order matters: this is the order written to Ecuser.txt and the
# Excel record. (label shown to user, field id written to file).
FIELDS = [
    ("ID", "id"),
    ("First Name", "First"),
    ("Last Name", "Last"),
    ("Pin", "Pin"),
    ("Passcode", "Passcode"),
]

OUTPUT_FILENAME = "Ecuser.txt"


def format_user_line(user):
    """Return a single Ecuser.txt line for one user dict.

    e.g. {"id": "1", "First": "John", ...} -> "id=1,First=John,..."
    """
    return ",".join("{}={}".format(field_id, user[field_id])
                    for _label, field_id in FIELDS)


class TimeClockApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Time Clock Bulk Upload")
        self.minsize(640, 520)

        # Users staged for this batch (list of dicts keyed by field id).
        self.users = []
        # Path to the linked Excel workbook (record of all users).
        self.workbook_path = None
        # All IDs we already know about: existing record + this batch.
        self.known_ids = set()

        self.entries = {}

        self._build_ui()

    # ----------------------------------------------------------------- UI
    def _build_ui(self):
        pad = {"padx": 8, "pady": 4}

        # --- Excel link row -------------------------------------------------
        link_frame = ttk.LabelFrame(self, text="Excel record (optional)")
        link_frame.pack(fill="x", **pad)

        self.link_var = tk.StringVar(value="No Excel file linked.")
        ttk.Label(link_frame, textvariable=self.link_var,
                  wraplength=420).grid(row=0, column=0, sticky="w", **pad)

        link_btns = ttk.Frame(link_frame)
        link_btns.grid(row=0, column=1, sticky="e", **pad)
        ttk.Button(link_btns, text="Link existing...",
                   command=self.link_existing_workbook).pack(side="left", padx=2)
        ttk.Button(link_btns, text="New record...",
                   command=self.create_new_workbook).pack(side="left", padx=2)
        link_frame.columnconfigure(0, weight=1)

        if not HAVE_OPENPYXL:
            self.link_var.set(
                "openpyxl not installed - Excel record feature disabled. "
                "Run 'pip install openpyxl' to enable it.")
            for child in link_btns.winfo_children():
                child.state(["disabled"])

        # --- Entry fields ---------------------------------------------------
        entry_frame = ttk.LabelFrame(self, text="New user")
        entry_frame.pack(fill="x", **pad)

        for row, (label, field_id) in enumerate(FIELDS):
            ttk.Label(entry_frame, text=label + ":").grid(
                row=row, column=0, sticky="e", **pad)
            entry = ttk.Entry(entry_frame, width=30)
            entry.grid(row=row, column=1, sticky="w", **pad)
            entry.bind("<Return>", self._on_return)
            self.entries[field_id] = entry
        entry_frame.columnconfigure(1, weight=1)

        ttk.Button(entry_frame, text="Add User",
                   command=self.add_user).grid(
            row=len(FIELDS), column=1, sticky="w", **pad)

        # --- Staged users list ---------------------------------------------
        list_frame = ttk.LabelFrame(self, text="Users in this batch")
        list_frame.pack(fill="both", expand=True, **pad)

        columns = [field_id for _label, field_id in FIELDS]
        self.tree = ttk.Treeview(list_frame, columns=columns, show="headings",
                                 height=8)
        for label, field_id in FIELDS:
            self.tree.heading(field_id, text=label)
            self.tree.column(field_id, width=110, anchor="w")
        self.tree.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)

        scroll = ttk.Scrollbar(list_frame, orient="vertical",
                               command=self.tree.yview)
        scroll.pack(side="left", fill="y", pady=8)
        self.tree.configure(yscrollcommand=scroll.set)

        ttk.Button(list_frame, text="Remove\nselected",
                   command=self.remove_selected).pack(side="left", padx=8, pady=8)

        # --- Bottom action bar ---------------------------------------------
        bottom = ttk.Frame(self)
        bottom.pack(fill="x", **pad)

        self.count_var = tk.StringVar(value="0 users staged")
        ttk.Label(bottom, textvariable=self.count_var).pack(side="left", padx=8)

        ttk.Button(bottom, text="Generate File",
                   command=self.generate_file).pack(side="right", padx=4)

        self.status_var = tk.StringVar(value="Ready.")
        status = ttk.Label(self, textvariable=self.status_var, relief="sunken",
                           anchor="w")
        status.pack(fill="x", side="bottom")

        # Focus first field
        self.entries[FIELDS[0][1]].focus_set()

    # ------------------------------------------------------------- helpers
    def _on_return(self, _event):
        self.add_user()

    def _set_status(self, msg):
        self.status_var.set(msg)

    def _refresh_count(self):
        self.count_var.set("{} user(s) staged".format(len(self.users)))

    # -------------------------------------------------------- Excel record
    def link_existing_workbook(self):
        path = filedialog.askopenfilename(
            title="Link existing Excel record",
            filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")])
        if not path:
            return
        try:
            ids = self._read_workbook_ids(path)
        except Exception as exc:  # noqa: BLE001 - report any read error
            messagebox.showerror("Could not read workbook", str(exc))
            return
        self.workbook_path = path
        # Reset known ids to existing record + ids already staged this batch.
        self.known_ids = set(ids) | {u["id"] for u in self.users}
        self.link_var.set("Linked: {}  ({} existing user(s))".format(
            os.path.basename(path), len(ids)))
        self._set_status("Linked workbook with {} existing users.".format(len(ids)))

    def create_new_workbook(self):
        path = filedialog.asksaveasfilename(
            title="Create new Excel record",
            defaultextension=".xlsx",
            initialfile="TimeClockUsers.xlsx",
            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Users"
            ws.append([label for label, _field_id in FIELDS])
            wb.save(path)
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Could not create workbook", str(exc))
            return
        self.workbook_path = path
        self.known_ids = {u["id"] for u in self.users}
        self.link_var.set("Linked: {}  (new, 0 existing user(s))".format(
            os.path.basename(path)))
        self._set_status("Created new workbook.")

    def _read_workbook_ids(self, path):
        """Return a set of existing IDs (as strings) from the workbook.

        Assumes the first column is the ID column and the first row is a
        header. Empty/blank IDs are ignored.
        """
        ids = set()
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:  # header
                continue
            if not row:
                continue
            value = row[0]
            if value is None:
                continue
            ids.add(self._normalize_id(value))
        wb.close()
        ids.discard("")
        return ids

    @staticmethod
    def _normalize_id(value):
        """Normalize an ID for comparison.

        Excel may return numeric IDs as floats (1.0). Treat 1, 1.0 and "1"
        as the same ID.
        """
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
        return text

    def _append_to_workbook(self, users):
        wb = load_workbook(self.workbook_path)
        ws = wb.active
        for user in users:
            ws.append([user[field_id] for _label, field_id in FIELDS])
        wb.save(self.workbook_path)

    # --------------------------------------------------------- user entry
    def add_user(self):
        user = {}
        for label, field_id in FIELDS:
            user[field_id] = self.entries[field_id].get().strip()

        # ID, First and Last are required. Pin/Passcode are optional.
        for label, field_id in FIELDS:
            if field_id in ("Pin", "Passcode"):
                continue
            if not user[field_id]:
                messagebox.showwarning("Missing field",
                                       "'{}' is required.".format(label))
                self.entries[field_id].focus_set()
                return

        # No commas or '=' in values - they would corrupt the file format.
        for label, field_id in FIELDS:
            if "," in user[field_id] or "=" in user[field_id]:
                messagebox.showwarning(
                    "Invalid character",
                    "'{}' cannot contain a comma or an equals sign.".format(label))
                self.entries[field_id].focus_set()
                return

        user_id = self._normalize_id(user["id"])
        user["id"] = user_id

        if user_id in self.known_ids:
            messagebox.showwarning(
                "Duplicate ID",
                "ID '{}' already exists. IDs must be unique.".format(user_id))
            self.entries["id"].focus_set()
            return

        self.users.append(user)
        self.known_ids.add(user_id)
        self.tree.insert("", "end",
                         values=[user[field_id] for _label, field_id in FIELDS])

        self._refresh_count()
        self._set_status("Added user {}.".format(user_id))

        # Clear fields and return focus to ID for the next entry.
        for _label, field_id in FIELDS:
            self.entries[field_id].delete(0, "end")
        self.entries["id"].focus_set()

    def remove_selected(self):
        selection = self.tree.selection()
        if not selection:
            return
        for item in selection:
            values = self.tree.item(item, "values")
            removed_id = self._normalize_id(values[0])
            # Remove the first matching staged user.
            for idx, user in enumerate(self.users):
                if user["id"] == removed_id:
                    del self.users[idx]
                    break
            self.known_ids.discard(removed_id)
            self.tree.delete(item)
        self._refresh_count()
        self._set_status("Removed selected user(s).")

    # ------------------------------------------------------- file output
    def generate_file(self):
        if not self.users:
            messagebox.showinfo("Nothing to generate",
                                "Add at least one user first.")
            return

        path = filedialog.asksaveasfilename(
            title="Save bulk upload file",
            defaultextension=".txt",
            initialfile=OUTPUT_FILENAME,
            filetypes=[("Text file", "*.txt"), ("All files", "*.*")])
        if not path:
            return

        try:
            with open(path, "w", encoding="utf-8", newline="\n") as fh:
                for user in self.users:
                    fh.write(format_user_line(user) + "\n")
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Could not write file", str(exc))
            return

        # Append this batch to the linked Excel record, if any.
        record_msg = ""
        if self.workbook_path and HAVE_OPENPYXL:
            try:
                self._append_to_workbook(self.users)
                record_msg = "\nAppended {} user(s) to {}.".format(
                    len(self.users), os.path.basename(self.workbook_path))
            except Exception as exc:  # noqa: BLE001
                messagebox.showerror(
                    "File written, but Excel update failed",
                    "{} was written, but the Excel record could not be "
                    "updated:\n{}".format(os.path.basename(path), exc))
                return

        messagebox.showinfo(
            "Done",
            "Wrote {} user(s) to {}.{}".format(
                len(self.users), os.path.basename(path), record_msg))
        self._set_status("Generated {} ({} users).".format(
            os.path.basename(path), len(self.users)))

        # Start a fresh batch but keep the workbook link and known IDs so
        # duplicates stay blocked across batches in the same session.
        self.users = []
        for item in self.tree.get_children():
            self.tree.delete(item)
        self._refresh_count()


def main():
    app = TimeClockApp()
    app.mainloop()


if __name__ == "__main__":
    main()
